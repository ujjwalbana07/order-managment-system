import base64
import re
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from xml.sax.saxutils import escape
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak
from core.templatetags.display import money, weight, percent

STYLES = getSampleStyleSheet()


def paragraph(value, style='Normal'):
    return Paragraph(escape(str(value)).replace('\n', '<br/>'), STYLES[style])


def table(rows, widths=None):
    rendered = [[paragraph(value) for value in row] for row in rows]
    result = Table(rendered, colWidths=widths, repeatRows=1, hAlign='LEFT')
    result.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e0eeee')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('LINEBELOW', (0, 0), (-1, -1), 0.3, colors.HexColor('#b8c8cc'))]))
    return result


def pdf_document(title, story, landscape_page=False):
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4) if landscape_page else A4,
        rightMargin=16*mm, leftMargin=16*mm, topMargin=16*mm, bottomMargin=18*mm, title=title)
    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont('Helvetica', 9)
        canvas.drawRightString(document.pagesize[0] - 16*mm, 10*mm, f'Page {document.page}')
        canvas.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def invoice_pdf(snapshot):
    issuer, recipient = snapshot['issuer'], snapshot['bill_to']
    story = [paragraph('Invoice ' + snapshot['invoice_no'], 'Title')]
    if snapshot.get('logo'):
        logo = Image(BytesIO(base64.b64decode(snapshot['logo'])))
        logo._restrictSize(45*mm, 20*mm)
        story.append(logo)
    story += [paragraph(issuer['company_name'], 'Heading2'), paragraph(issuer['address']),
        paragraph(f"{issuer['phone']} · {issuer['email']} · {issuer['tax_id']}"),
        paragraph('Bill to: ' + recipient['billing_name'], 'Heading2'), paragraph(recipient['billing_address']),
        paragraph(f"{recipient['email']} · {recipient['phone']} · {recipient['tax_id']}"), Spacer(1, 6*mm)]
    order = snapshot['order']
    story += [paragraph(f"Issued {datetime.fromisoformat(snapshot['issued_at']).strftime('%d-%b-%Y')} · Sales {order['sales_no']} · SR {order['serial_no']}"),
        paragraph(f"Challan {order['challan_no']} · Manufacturing {order['mfg_order_no']}"),
        paragraph(f"{order['item_title']} · Quantity {order['quantity']}"), Spacer(1, 5*mm)]
    rows = [['Component', 'Weight', 'Rate (INR)', 'Amount (INR)']]
    rows += [[line['component'], weight(line['weight']), money(line['rate']) if line['rate'] else '', money(line['amount'])] for line in snapshot['lines']]
    story.append(table(rows, [45*mm]*4))
    for label, field in [('Total bill', 'total_bill'), ('Paid out to date' if snapshot.get('payment_direction', 'paid_out') == 'paid_out' else 'Received to date', 'total_received'), ('Outstanding', 'total_outstanding')]:
        story.append(paragraph(f"{label}: INR {money(snapshot['totals'][field])}", 'Heading3'))
    if Decimal(snapshot.get('gst_rate_percent', '0')):
        story.append(paragraph(f"GST ({snapshot['gst_rate_percent']}%): INR {money(snapshot['gst_amount'])}"))
    if snapshot.get('photo'):
        image = Image(BytesIO(base64.b64decode(snapshot['photo'])))
        image._restrictSize(40*mm, 40*mm)
        story.append(image)
    story.append(paragraph(issuer['invoice_footer_text']))
    return pdf_document(snapshot['invoice_no'], story)


def receipt_pdf(payment, balance, direction='paid_out'):
    order = payment.order
    title = 'Payment receipt' if direction == 'paid_out' else 'Receipt'
    story = [paragraph(title + ' ' + payment.receipt_no, 'Title'),
        paragraph(f"{payment.status} · {payment.payment_date:%d-%b-%Y}", 'Heading2'),
        paragraph(f"{order.account.billing_name} · Sales {order.sales_no} · SR {order.serial_no}"),
        paragraph(f"Challan {order.challan_no} · Manufacturing {order.mfg_order_no}"),
        paragraph(order.item_title), paragraph(f"{payment.method} · Reference {payment.reference_no}"),
        paragraph(f"Amount: INR {money(payment.amount)}", 'Heading2')]
    amounts = {row.component: row.amount for row in payment.allocations.all()}
    story.append(table([['Component', 'This payment', 'Remaining after this payment']] +
        [[row['name'], money(amounts.get(row['name'], 0)), money(row['outstanding'])] for row in balance['components']], [60*mm]*3))
    story += [paragraph(f"Remaining total: INR {money(balance['total_outstanding'])}", 'Heading2'),
        paragraph(f"Recorded by {payment.created_by.email}"), paragraph(payment.remarks)]
    if payment.status == 'Voided':
        story.append(paragraph('Voided: ' + payment.void_reason, 'Heading2'))
    return pdf_document(payment.receipt_no, story)


def safe_cell(value):
    # Keep external identifiers as literal text, never executable spreadsheet formulas.
    if isinstance(value, str):
        value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value)
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def workbook_bytes(sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, headers, rows, totals in sheets:
        sheet = workbook.create_sheet(name[:31])
        sheet.append(headers)
        for row in rows:
            sheet.append([safe_cell(value) for value in row])
        if totals:
            sheet.append(totals)
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = f'A1:{sheet.cell(max(1, sheet.max_row - bool(totals)), len(headers)).coordinate}'
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, Decimal):
                    label = headers[cell.column-1].lower()
                    cell.number_format = '0.00%' if 'purity' in label else '0.0000' if 'fx' in label else '0.000' if any(word in label for word in ('weight', ' wt', 'carat', 'pure 995')) else '#,##0.00'
                elif hasattr(cell.value, 'year'):
                    cell.number_format = 'dd-mmm-yyyy'
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(50, max(15, len(str(column[0].value)) + 2))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def report_pdf(title, headers, rows, filter_text, generated_at):
    story = [paragraph(title, 'Title'), paragraph('Generated: ' + generated_at), paragraph('Filters: ' + filter_text), Spacer(1, 5*mm)]
    # Wide exports are split into panels, retaining the first identity column.
    for start in range(1, len(headers), 6):
        indices = [0] + list(range(start, min(start + 6, len(headers))))
        if start > 1:
            story.append(PageBreak())
            story += [paragraph(title + ' (continued)', 'Heading1'), paragraph('Generated: ' + generated_at), paragraph('Filters: ' + filter_text)]
        formatted = [[headers[index] for index in indices]]
        for row in rows:
            values = []
            for index in indices:
                value = row[index]
                label = headers[index].lower()
                if isinstance(value, Decimal):
                    value = percent(value) if 'purity' in label else format(value, '.4f') if 'fx' in label else weight(value) if any(word in label for word in ('weight', 'carat', 'pure 995')) else money(value)
                elif hasattr(value, 'strftime'):
                    value = value.strftime('%d-%b-%Y')
                values.append(value)
            formatted.append(values)
        story.append(table(formatted, [265*mm/len(indices)]*len(indices)))
    return pdf_document(title, story, landscape_page=True)
