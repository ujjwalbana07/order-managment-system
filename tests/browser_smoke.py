"""Optional end-to-end smoke test against an isolated running test instance.

OMS_SMOKE_CREDENTIALS points to a private JSON file containing email, password,
and account_id. OMS_SMOKE_URL defaults to http://127.0.0.1:8765.
"""
from datetime import date
from io import BytesIO
import json
import os
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
from playwright.sync_api import sync_playwright


def main():
    credentials = json.loads(Path(os.environ['OMS_SMOKE_CREDENTIALS']).read_text())
    root = Path(__file__).resolve().parents[1] / 'artifacts'
    root.mkdir(exist_ok=True)
    url = os.environ.get('OMS_SMOKE_URL', 'http://127.0.0.1:8765')
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(channel='chrome', headless=True)
        first_sales = int(os.environ.get('OMS_SMOKE_FIRST_SALES', '501'))
        for label, viewport, sales_no in [('desktop', {'width': 1440, 'height': 1000}, str(first_sales)), ('mobile', {'width': 390, 'height': 844}, str(first_sales + 1))]:
            context = browser.new_context(viewport=viewport, base_url=url, accept_downloads=True)
            page = context.new_page()
            errors = []
            page.on('pageerror', lambda error: errors.append(str(error)))
            page.goto('/login/')
            page.locator('#id_username').fill(credentials['email'])
            page.locator('#id_password').fill(credentials['password'])
            page.get_by_role('button', name='Sign in', exact=True).click()
            page.wait_for_url(url + '/')
            if not page.locator('.account-picker').get_attribute('open') and not page.locator('#account-select').is_visible():
                page.locator('.account-picker > summary').click()
            page.locator('#account-select').select_option(str(credentials['account_id']))
            page.get_by_role('button', name='Switch account', exact=True).click()
            page.goto('/orders/new/')
            values = {'sales_no': sales_no, 'order_date': date.today().isoformat(), 'buyer_username': 'smoke-buyer',
                'item_title': 'Golden test ring', 'quantity': '1', 'usd_sold': '1859.26', 'fx_rate': '91.0000',
                'inr_sold': '169192.66', 'ship_charges': '5000', 'platform_fees_inr': '0', 'purity': '59.5',
                'gross_wt': '4.24', 'dia_ct': '7.56', 'other_wt': '0', 'gold_rate': '15000', 'lab_rate': '1250', 'diamond_value': '87740'}
            for name, value in values.items():
                page.locator('#id_' + name).fill(value)
            page.locator('#id_image').set_input_files(Path(__file__).resolve().parents[1] / 'oms-pack/seed/images/101.png')
            page.wait_for_function("document.querySelector('#cost-preview').textContent.includes('115619.75')")
            page.get_by_role('button', name='Save order', exact=True).click()
            page.wait_for_url('**/orders/*/')
            assert '/new/' not in page.url
            detail_url = page.url
            assert '1,15,619.75' in page.locator('main').inner_text()
            page.get_by_role('link', name='Record payment', exact=True).click()
            page.locator('#id_amount').fill('10000.00')
            page.get_by_role('button', name='Split automatically', exact=True).click()
            page.wait_for_function("document.querySelector('#id_Gold').value === '10000.00'")
            page.get_by_role('button', name='Save payment', exact=True).click()
            page.wait_for_url(detail_url)
            assert 'Partially Paid' in page.locator('main').inner_text()
            with page.expect_download() as invoice_download:
                page.get_by_role('button', name='Issue invoice', exact=True).click()
            invoice_path = root / f'{label}-invoice.pdf'
            invoice_download.value.save_as(invoice_path)
            assert 'Smoke Jewels' in '\n'.join(p.extract_text() for p in PdfReader(invoice_path).pages)
            page.goto(detail_url)
            with page.expect_download() as receipt_download:
                page.locator('a[href$="/receipt/"]').click()
            receipt_path = root / f'{label}-receipt.pdf'
            receipt_download.value.save_as(receipt_path)
            assert '10,000.00' in '\n'.join(p.extract_text() for p in PdfReader(receipt_path).pages)
            page.screenshot(path=str(root / f'{label}-order.png'), full_page=True)
            page.goto('/orders/?q=' + sales_no)
            with page.expect_download() as excel_download:
                page.get_by_role('link', name='Export Excel', exact=True).click()
            workbook_path = root / f'{label}-orders.xlsx'
            excel_download.value.save_as(workbook_path)
            assert load_workbook(workbook_path).active.max_row == 3
            if label == 'mobile':
                assert page.locator('.order-cards').is_visible()
                assert not page.locator('.order-grid').is_visible()
            assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth'), 'Page overflows viewport'
            page.screenshot(path=str(root / f'{label}-orders.png'), full_page=True)
            assert not errors, errors
            print(f'{label} {viewport["width"]}px: login, order, image, costing preview, payment split, invoice, receipt, Excel, layout passed')
            context.close()
        browser.close()


if __name__ == '__main__':
    main()
