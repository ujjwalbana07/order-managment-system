from datetime import timedelta
from decimal import Decimal
from io import BytesIO
import pytest
from openpyxl import load_workbook
from django.core.exceptions import PermissionDenied
from django.utils import timezone
from accounts.models import User, EbayAccount
from core.models import Settings
from core.services import save_record
from documents.services import issue_invoice
from documents.renderers import invoice_pdf
from orders.services import save_order
from payments.services import post_payment, balances

pytestmark = pytest.mark.django_db


@pytest.fixture
def issuer(actor):
    return save_record(Settings(company_name='Jewels', address='Main street', phone='123', email='issuer@example.com'), actor=actor)


def test_invoice_snapshot_revised_number(make_order, actor, issuer):
    order = save_record(make_order(), actor=actor)
    invoice = issue_invoice(actor=actor, order_id=order.pk)
    assert invoice.invoice_no.startswith(f'INV-{timezone.localdate().year}-')
    assert invoice.snapshot['totals']['total_bill'] == '115619.75'
    original = invoice.snapshot.copy()
    save_order(actor=actor, order_id=order.pk, version=1, data={'gold_rate': Decimal('16000')})
    issuer.company_name = 'Changed company'
    save_record(issuer, actor=actor)
    revised = issue_invoice(actor=actor, order_id=order.pk)
    invoice.refresh_from_db()
    assert invoice.status == 'Superseded' and invoice.snapshot == original
    assert revised.invoice_no != invoice.invoice_no
    assert revised.snapshot['issuer']['company_name'] == 'Changed company'
    assert invoice_pdf(invoice.snapshot).startswith(b'%PDF-')


def test_receipt_chronology(make_order, actor, client):
    order = save_record(make_order(), actor=actor)
    later = post_payment(actor=actor, order_id=order.pk, amount=Decimal('10'), method='Cash', payment_date=timezone.localdate(), allocations={'Gold': Decimal('10')})
    earlier = post_payment(actor=actor, order_id=order.pk, amount=Decimal('5'), method='Cash', payment_date=timezone.localdate()-timedelta(days=1), allocations={'Gold': Decimal('5')})
    assert balances(order, through=earlier)['total_received'] == Decimal('5')
    assert balances(order, through=later)['total_received'] == Decimal('15')
    client.force_login(actor)
    response = client.get(f'/payments/{later.pk}/receipt/')
    assert response.status_code == 200 and response.content.startswith(b'%PDF-')


def test_filtered_exports_and_all_pdf_variants(make_order, actor, issuer, client):
    order = save_record(make_order(), actor=actor)
    save_record(make_order(sales_no=102), actor=actor)
    client.force_login(actor)
    response = client.get('/reports/orders.xlsx', {'q': '101'})
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    assert sheet.freeze_panes == 'A2'
    assert sheet.max_row == 3
    assert sheet['C2'].value == 101
    assert Decimal(str(sheet['X2'].value)) == order.total_bill
    assert sheet['X2'].number_format == '#,##0.00'
    for kind in ('orders', 'payments', 'outstanding', 'statement'):
        response = client.get(f'/reports/{kind}.pdf')
        assert response.status_code == 200 and response.content.startswith(b'%PDF-')
    for url in ('/reports/', '/manage/settings/', '/'):
        assert client.get(url).status_code == 200


def test_document_access_and_formula_injection(make_order, actor, issuer, client, account):
    order = save_record(make_order(item_title='=DANGEROUS()'), actor=actor)
    invoice = issue_invoice(actor=actor, order_id=order.pk)
    outsider = User.objects.create_user(email='outsider@example.com', role='Accounts')
    client.force_login(outsider)
    assert client.get(f'/invoices/{invoice.pk}/download/').status_code == 200
    sheet = load_workbook(BytesIO(client.get('/reports/orders.xlsx').content)).active
    dangerous = [cell for cell in sheet['G'] if cell.value and str(cell.value).startswith("'=")]
    assert dangerous and dangerous[0].data_type == 's'
    outsider.role = 'Staff'
    outsider.save()
    assert client.get('/reports/orders.xlsx').status_code == 403
    assert client.post(f'/orders/{order.pk}/invoice/').status_code == 403


def test_payment_filters_and_totals(make_order, actor, client):
    order = save_record(make_order(), actor=actor)
    save_record(make_order(sales_no=102), actor=actor)
    post_payment(actor=actor, order_id=order.pk, amount=Decimal('5'), method='Cash', payment_date=timezone.localdate(), allocations={'Gold': Decimal('5')})
    client.force_login(actor)
    response = client.get('/orders/', {'payment_status': 'Partially Paid'})
    assert response.context['page'].paginator.count == 1
    assert response.context['totals']['received'] == Decimal('5')
    assert response.context['totals']['outstanding'] == order.total_bill - Decimal('5')


def test_pdf_pages_parse_and_snapshot_text(make_order, actor, issuer):
    PdfReader = pytest.importorskip('pypdf').PdfReader
    order = save_record(make_order(), actor=actor)
    invoice = issue_invoice(actor=actor, order_id=order.pk)
    reader = PdfReader(BytesIO(invoice_pdf(invoice.snapshot)))
    text = '\n'.join(page.extract_text() for page in reader.pages)
    assert 'Jewels' in text and '1,15,619.75' in text
    assert invoice.invoice_no in text and 'Page 1' in text
