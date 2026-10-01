from io import BytesIO, StringIO
import json
import os
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
import sqlite3
import pytest
from django.core.management import call_command
from django.db import DatabaseError
from accounts.models import User
from audit.models import AuditLog
from core.backups import full_backup
from core.services import save_record

pytestmark = pytest.mark.django_db


def test_health_and_request_ids(client):
    response = client.get('/healthz')
    assert response.status_code == 200 and response.content == b'OK'
    assert len(response['X-Request-ID']) == 32
    with patch('core.views.connection.cursor', side_effect=DatabaseError('Do not expose this')):
        response = client.get('/healthz')
    assert response.status_code == 503 and b'Do not expose' not in response.content


def test_friendly_error_pages(client):
    assert client.get('/not-a-real-page/').status_code == 404
    from core.views import server_error
    from django.test import RequestFactory
    request = RequestFactory().get('/')
    request.request_id = 'reference123'
    response = server_error(request)
    assert response.status_code == 500 and b'reference123' in response.content


def test_backup_contains_deleted_orders_and_photos(make_order, actor, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    (tmp_path / 'photo.jpg').write_bytes(b'photo fixture')
    order = save_record(make_order(is_deleted=True), actor=actor)
    content = full_backup(actor=actor)
    with ZipFile(BytesIO(content)) as archive:
        assert 'data.xlsx' in archive.namelist()
        assert archive.read('media/photo.jpg') == b'photo fixture'
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(archive.read('data.xlsx')))
        assert book['orders.Order'].max_row == 2
        assert book['orders.Order']['A2'].value == order.pk
        assert 'password' not in [cell.value for cell in book['accounts.User'][1]]
    assert AuditLog.objects.filter(action='backup_download', actor=actor).count() == 1


def test_audit_and_backup_owner_only(client, actor):
    client.force_login(actor)
    assert client.get('/manage/audit/').status_code == 200
    assert client.get('/manage/backups/').status_code == 200
    assert client.post('/manage/backups/').status_code == 200
    user = User.objects.create_user(email='noaccess@example.com', role='Accounts')
    client.force_login(user)
    for path in ('/manage/backups/', '/manage/audit/'):
        assert client.get(path).status_code == 403
        assert client.post(path).status_code == 403


@pytest.mark.django_db(transaction=True)
def test_backup_command_restorable_sqlite_and_retention(actor, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path / 'media'
    root = tmp_path / 'backups'
    root.mkdir()
    old = root / 'oms_backup_old.zip'
    old.write_bytes(b'old')
    os.utime(old, (1, 1))
    unrelated = root / 'unrelated.zip'
    unrelated.write_bytes(b'keep')
    output = StringIO()
    call_command('backup_db', directory=str(root), stdout=output)
    backups = list(root.glob('oms_backup_*.zip'))
    assert len(backups) == 1 and unrelated.exists() and not old.exists()
    with ZipFile(backups[0]) as archive:
        assert json.loads(archive.read('metadata.json'))['database_vendor'] == 'sqlite'
        database = tmp_path / 'restored.sqlite3'
        database.write_bytes(archive.read('database.sqlite3'))
    restored = sqlite3.connect(database)
    try:
        assert restored.execute('PRAGMA integrity_check').fetchone() == ('ok',)
        assert restored.execute('SELECT email FROM accounts_user WHERE id = ?', (actor.pk,)).fetchone() == (actor.email,)
    finally:
        restored.close()
