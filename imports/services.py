from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
import re
from openpyxl import load_workbook
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import transaction
from django.shortcuts import get_object_or_404
from django.utils import timezone
from accounts.access import require
from accounts.models import EbayAccount
from audit.models import AuditLog
from core.models import default_fx_rate, default_gold_rate
from core.transactions import retry_locked
from imports.models import ImportBatch
from orders.images import prepare_image
from orders.models import Order
from orders.services import save_order

HEADERS = {
    'Sales No.': 'sales_no', 'Order Date': 'order_date', 'Ship By Date': 'ship_by_date', 'Buyer Username': 'buyer_username',
    'Photos': 'image', 'Item Title': 'item_title', 'Quantity': 'quantity', '$ me Sold For': 'usd_sold', 'INR RS.': 'inr_sold',
    'PURITY': 'purity', 'GROSS WT': 'gross_wt', 'TTL DIA WT CTS.': 'dia_ct', 'OTHER': 'other_wt', 'NET WT': 'net_wt',
    'PURE 995': 'pure_995', 'IGI': 'igi_cert_no', 'LAB RATE': 'lab_rate', 'LAB AMOUNT': 'labour_amount',
    'DIAMOND VALUE': 'diamond_value', 'AMOUNT': 'total_bill', 'Gold Ret': 'gold_amount', 'Tracking No.': 'tracking_no',
    'Ship Charges': 'ship_charges', 'NET Earnings': 'net_earnings',
}
DERIVED = {'net_wt', 'pure_995', 'labour_amount', 'total_bill', 'gold_amount', 'net_earnings'}
REQUIRED = {'sales_no', 'order_date', 'buyer_username', 'item_title', 'quantity', 'inr_sold', 'purity', 'gross_wt', 'dia_ct', 'other_wt', 'lab_rate', 'diamond_value', 'ship_charges'}


def normalize(value):
    return re.sub(r'[^a-z0-9]', '', str(value or '').casefold())


def decimal_value(value, optional=False):
    if optional and (value is None or str(value).strip() == ''):
        return None
    try:
        result = Decimal(str(value).replace('$', '').replace(',', '').strip())
    except InvalidOperation as exc:
        raise ValidationError('Enter a valid number in the spreadsheet.') from exc
    if not result.is_finite() or (result and abs(result.adjusted()) > 18):
        raise ValidationError('Spreadsheet numbers must be finite.')
    return result


def date_value(value, optional=False):
    if optional and not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for pattern in ('%Y-%m-%d', '%d-%b-%Y', '%d/%m/%Y'):
        try:
            return datetime.strptime(str(value).strip(), pattern).date()
        except ValueError:
            continue
    raise ValidationError('Use a real Excel date or YYYY-MM-DD date.')


def open_workbook(content):
    if len(content) > 10 * 1024 * 1024:
        raise ValidationError('Choose an Excel file smaller than 10 MB.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise ValidationError('This workbook is too large when opened. Split it into smaller files.')
        return load_workbook(BytesIO(content), data_only=True)
    except (BadZipFile, KeyError, ValueError, OSError, ParseError) as exc:
        raise ValidationError('Choose a valid .xlsx workbook.') from exc


def serializable_preview(result):
    return {
        'rows': [{'row': row['row'], 'sales_no': row['sales_no'], 'status': row['status'], 'messages': row['messages']} for row in result['rows']],
        'gaps': [[start, end] for start, end in result['gaps']],
        'has_errors': result['has_errors'],
    }


def preview(content, *, account, actor, gold_rate=None):
    require(actor, 'import')
    if not EbayAccount.objects.for_user(actor).filter(pk=account.pk, active=True).exists():
        raise ValidationError('Select an active account you have access to.')
    book = open_workbook(content)
    sheet = book.active
    if sheet.max_row > 10001 or sheet.max_column > 100:
        raise ValidationError('Use at most 10,000 rows and 100 columns per import.')
    mapping = {normalize(label): field for label, field in HEADERS.items()}
    columns = {}
    for index, cell in enumerate(sheet[1]):
        field = mapping.get(normalize(cell.value))
        if field:
            if field in columns:
                raise ValidationError(f'The spreadsheet repeats the {field} column.')
            columns[field] = index
    if REQUIRED - set(columns):
        raise ValidationError('Missing spreadsheet columns: ' + ', '.join(sorted(REQUIRED - set(columns))))
    images = {}
    for image in sheet._images:
        if hasattr(image.anchor, '_from'):
            images.setdefault(image.anchor._from.row + 1, []).append(image._data())
    existing = set(Order.objects.filter(account=account).values_list('sales_no', flat=True))
    seen = set()
    rows, sales = [], []
    for number, cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if all(value is None for value in cells):
            continue
        row_images = images.get(number, [])
        row = {'row': number, 'sales_no': None, 'status': 'New', 'messages': [], 'data': {}, 'image_bytes': row_images[0] if row_images else None,
            'image2_bytes': row_images[1] if len(row_images) > 1 else None}
        try:
            values = {name: cells[index] for name, index in columns.items()}
            sales_number = decimal_value(values['sales_no'])
            if sales_number != sales_number.to_integral_value():
                raise ValidationError('Sales number must be an integer.')
            row['sales_no'] = int(sales_number)
            sales.append(int(sales_number))
            if row['sales_no'] in existing or row['sales_no'] in seen:
                row['status'] = 'Duplicate'
                row['messages'].append('This sales number already exists in this account or file and will be skipped.')
                rows.append(row)
                continue
            seen.add(row['sales_no'])
            data = {'account': account, 'sales_no': row['sales_no'], 'order_date': date_value(values['order_date']),
                'ship_by_date': date_value(values.get('ship_by_date'), optional=True),
                'gold_rate': decimal_value(values.get('gold_rate'), optional=True) or gold_rate or default_gold_rate(),
                'fx_rate': default_fx_rate(), 'platform_fees_inr': Decimal('0')}
            for name in ('buyer_username', 'item_title', 'igi_cert_no', 'tracking_no'):
                data[name] = str(values.get(name) or '').strip()
            for name in ('usd_sold', 'inr_sold', 'purity', 'gross_wt', 'dia_ct', 'other_wt', 'lab_rate', 'diamond_value', 'ship_charges'):
                value = decimal_value(values.get(name), optional=name in {'usd_sold', 'other_wt'})
                data[name] = value if value is not None else Decimal('0')
            if data['purity'] > 1:
                data['purity'] /= Decimal('100')
            quantity = decimal_value(values['quantity'])
            if quantity != quantity.to_integral_value():
                raise ValidationError('Quantity must be an integer.')
            data['quantity'] = int(quantity)
            order = Order(**data, created_by=actor, updated_by=actor)
            order.full_clean()
            for name in DERIVED:
                if values.get(name) is not None:
                    value = decimal_value(values[name])
                    expected = getattr(order, name)
                    if value != expected:
                        row['messages'].append(f'{name}: sheet {value}, system {expected}')
            if row['image_bytes']:
                prepare_image(SimpleUploadedFile('photo', row['image_bytes']))
            if row['image2_bytes']:
                prepare_image(SimpleUploadedFile('photo2', row['image2_bytes']))
            row['data'] = data
            row['status'] = 'Difference' if row['messages'] else 'New'
        except (ValidationError, ValueError, TypeError, InvalidOperation) as exc:
            row['status'] = 'Error'
            row['messages'] = exc.messages if isinstance(exc, ValidationError) else ['This row contains an invalid value. Check numbers and dates.']
        rows.append(row)
    sorted_sales = sorted(set(sales))
    gaps = [(a + 1, b - 1) for a, b in zip(sorted_sales, sorted_sales[1:]) if b > a + 1]
    return {'rows': rows, 'gaps': gaps, 'has_errors': any(row['status'] == 'Error' for row in rows)}


@transaction.atomic
def create_batch(*, actor, account, upload, gold_rate=None):
    content = upload.read(10 * 1024 * 1024 + 1)
    result = preview(content, account=account, actor=actor, gold_rate=gold_rate)
    batch = ImportBatch(actor=actor, account=account, filename=Path(upload.name).name, source=content,
        preview_cache=serializable_preview(result), gold_rate=gold_rate or default_gold_rate())
    batch.save()
    AuditLog.objects.create(actor=actor, action='import_preview', object_type='imports.ImportBatch', object_id=str(batch.pk),
        account=account, after={'filename': batch.filename})
    return batch


@retry_locked
@transaction.atomic
def commit_batch(*, actor, batch_id):
    require(actor, 'import')
    batch = get_object_or_404(ImportBatch.objects.select_for_update(), pk=batch_id, actor=actor,
        account__in=EbayAccount.objects.for_user(actor))
    if batch.committed_at:
        return 0
    # One account lock serialises imports targeting the same sales-number namespace.
    account = EbayAccount.objects.select_for_update().get(pk=batch.account_id)
    result = preview(bytes(batch.source), account=account, actor=actor, gold_rate=batch.gold_rate)
    if result['has_errors']:
        raise ValidationError('Fix every Error row before importing. No orders were created.')
    created = []
    try:
        for row in result['rows']:
            if row['status'] == 'Duplicate':
                continue
            image = SimpleUploadedFile('photo', row['image_bytes']) if row['image_bytes'] else None
            image2 = SimpleUploadedFile('photo2', row['image2_bytes']) if row['image2_bytes'] else None
            order = save_order(actor=actor, data=row['data'], image=image, image2=image2)
            created.append(order)
        batch.committed_at = timezone.now()
        batch.created_count = len(created)
        batch.save()
        AuditLog.objects.create(actor=actor, action='import', object_type='imports.ImportBatch', object_id=str(batch.pk),
            account=account, after={'orders': [order.pk for order in created], 'created_count': len(created)})
    except Exception:
        for order in created:
            for file in (order.image, order.thumbnail, order.image2, order.thumbnail2):
                if file:
                    file.storage.delete(file.name)
        raise
    return len(created)
